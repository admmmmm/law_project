import csv
import io
import json


def parse_structured_content(
    source_type: str,
    content: str,
    raw_bytes: bytes | None = None,
) -> list[dict[str, str]]:
    """
    解析结构化内容并统一输出为行字典列表。
    """
    normalized = source_type.lower()

    if normalized == "csv":
        reader = csv.DictReader(io.StringIO(content))
        return [{k: str(v) for k, v in row.items()} for row in reader]

    if normalized == "json":
        data = json.loads(content)
        if isinstance(data, dict):
            return [{str(k): str(v) for k, v in data.items()}]
        if isinstance(data, list):
            rows: list[dict[str, str]] = []
            for item in data:
                if isinstance(item, dict):
                    rows.append({str(k): str(v) for k, v in item.items()})
            return rows
        raise ValueError("JSON 内容格式不支持，需为对象或对象数组。")

    if normalized == "流水":
        rows: list[dict[str, str]] = []
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        for line in lines:
            if "，" in line:
                parts = [p.strip() for p in line.split("，") if p.strip()]
            else:
                parts = [p.strip() for p in line.split(",") if p.strip()]
            if len(parts) >= 3:
                rows.append({"subject": parts[0], "relation": parts[1], "object": parts[2]})
        if rows:
            return rows
        raise ValueError("流水文本需按“主体，关系，对象”格式提供。")

    if normalized == "xlsx":
        if raw_bytes is None:
            raise ValueError("XLSX 文件为空，无法解析。")
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise ValueError("解析 XLSX 需要安装 openpyxl。") from exc

        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        headers = next(rows_iter, None)
        if not headers:
            return []

        normalized_headers = [str(h).strip() if h is not None else "" for h in headers]
        rows: list[dict[str, str]] = []
        for row in rows_iter:
            row_dict: dict[str, str] = {}
            for key, value in zip(normalized_headers, row):
                if key:
                    row_dict[key] = "" if value is None else str(value)
            rows.append(row_dict)
        return rows

    raise ValueError(f"不支持的结构化类型: {source_type}")
