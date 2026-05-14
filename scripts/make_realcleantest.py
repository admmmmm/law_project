from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "realtest"
TARGET = ROOT / "data" / "realcleantest"

REPLACEMENTS = {
    "杨周武": "杨承泽",
    "王静": "王雨晴",
    "何晓初": "何婉初",
    "刘力飚": "刘建峰",
    "罗贤涛": "罗启涛",
    "易承桂": "易景文",
    "江军": "江立",
    "汪春蓉": "汪雅蓉",
    "赵志高": "赵文高",
    "张夏天": "张明夏",
    "罗宇": "罗成",
    "陈三一": "陈志一",
    "曾黎": "曾岚",
    "饶蝶": "饶雯",
    "周芷": "周晴",
    "张三": "张远",
    "李四": "李成",
    "赵五": "赵明",
}


def replace_text(text: str) -> str:
    for old, new in REPLACEMENTS.items():
        text = text.replace(old, new)
    return text


def rewrite_text_file(path: Path) -> None:
    text = path.read_text(encoding="utf-8")
    path.write_text(replace_text(text), encoding="utf-8", newline="")


def rewrite_xlsx(path: Path) -> None:
    wb = load_workbook(path)
    changed = False
    for ws in wb.worksheets:
        ws.title = replace_text(ws.title)[:31]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    new_value = replace_text(cell.value)
                    if new_value != cell.value:
                        cell.value = new_value
                        changed = True
    if changed:
        wb.save(path)


def rename_path(path: Path) -> None:
    new_name = replace_text(path.name)
    if new_name == path.name:
        return
    target = path.with_name(new_name)
    if target.exists():
        raise FileExistsError(f"target exists: {target}")
    path.rename(target)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if TARGET.exists():
        resolved = TARGET.resolve()
        if ROOT.resolve() not in resolved.parents:
            raise RuntimeError(f"refusing to remove outside project: {resolved}")
        shutil.rmtree(TARGET)
    shutil.copytree(SOURCE, TARGET)

    for path in TARGET.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() in {".md", ".txt", ".csv", ".json"}:
            rewrite_text_file(path)
        elif path.suffix.lower() == ".xlsx":
            rewrite_xlsx(path)

    for path in sorted(TARGET.rglob("*"), key=lambda item: len(item.parts), reverse=True):
        rename_path(path)

    print(f"created {TARGET}")
    print("replacement map:")
    for old, new in REPLACEMENTS.items():
        print(f"  {old} -> {new}")


if __name__ == "__main__":
    main()
