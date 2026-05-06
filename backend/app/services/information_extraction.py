from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from app.schemas.ingestion import ExtractedTriple, ExtractionResult, PassageRecord

STRUCTURED_TYPES = {"csv", "xlsx", "json", "流水", "bank_flow", "transaction"}


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def resolve_source_type(source_type: str | None, filename: str | None = None) -> str:
    value = (source_type or "").strip()
    if value and value != "unknown":
        return value
    if filename:
        suffix = Path(filename).suffix.lower().lstrip(".")
        if suffix:
            return suffix
    return "unknown"


def extract_content(source_type: str, raw_bytes: bytes | None = None, content: str | None = None) -> str:
    normalized = source_type.lower()
    if raw_bytes is None:
        return content or ""
    if normalized in {"txt", "text", "csv", "json"}:
        return decode_text(raw_bytes)
    if normalized == "xlsx":
        return content or ""
    if normalized == "doc":
        raise ValueError("暂不支持直接解析 .doc，请转成 .docx、.txt、.pdf、.csv 或 .xlsx 后再上传。")
    if normalized == "docx":
        try:
            from docx import Document
        except Exception as exc:
            return extract_docx_text(raw_bytes)
        doc = Document(io.BytesIO(raw_bytes))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    if normalized == "pdf":
        try:
            from pypdf import PdfReader
        except Exception as exc:
            raise ValueError("解析 PDF 需要安装 pypdf。") from exc
        reader = PdfReader(io.BytesIO(raw_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    return decode_text(raw_bytes)


def extract_docx_text(raw_bytes: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(raw_bytes)) as archive:
            document_xml = archive.read("word/document.xml")
    except Exception as exc:
        raise ValueError("无法解析 DOCX 文档正文。") from exc

    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = ElementTree.fromstring(document_xml)
    paragraphs: list[str] = []
    for paragraph in root.findall(".//w:p", namespace):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace)).strip()
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)


def route_extraction(
    *,
    source_type: str,
    content: str,
    raw_bytes: bytes | None = None,
    filename: str | None = None,
    evidence_id: str | None = None,
) -> ExtractionResult:
    resolved = resolve_source_type(source_type, filename)
    if resolved.lower() in STRUCTURED_TYPES or resolved in STRUCTURED_TYPES:
        rows = parse_structured_content(source_type=resolved, content=content, raw_bytes=raw_bytes)
        triples = map_rows_to_triples(rows, evidence_id=evidence_id, filename=filename)
        return ExtractionResult(
            route="structured",
            triples=triples,
            passages=build_passages_from_triples(triples),
            metadata={"source_type": resolved, "rows_count": len(rows)},
        )
    triples, passages = extract_unstructured_triples(content, evidence_id=evidence_id)
    return ExtractionResult(
        route="unstructured",
        triples=triples,
        passages=passages,
        metadata={"source_type": resolved, "chars": len(content)},
    )


def parse_structured_content(source_type: str, content: str, raw_bytes: bytes | None = None) -> list[dict[str, str]]:
    normalized = source_type.lower()
    if normalized == "json":
        data = json.loads(content)
        if isinstance(data, dict):
            return [{str(k): _cell(v) for k, v in data.items()}]
        if isinstance(data, list):
            return [{str(k): _cell(v) for k, v in row.items()} for row in data if isinstance(row, dict)]
        raise ValueError("JSON 内容必须是对象或对象数组。")
    if normalized == "xlsx":
        if raw_bytes is None:
            raise ValueError("解析 XLSX 需要原始文件字节。")
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise ValueError("解析 XLSX 需要安装 openpyxl。") from exc
        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
        rows: list[dict[str, str]] = []
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, None)
            if not headers:
                continue
            keys = [_cell(h) for h in headers]
            for values in iterator:
                row = {key: _cell(value) for key, value in zip(keys, values) if key}
                if any(row.values()):
                    row["_sheet"] = sheet.title
                    rows.append(row)
        return rows
    reader = csv.DictReader(io.StringIO(content))
    if reader.fieldnames:
        return [{str(k): _cell(v) for k, v in row.items()} for row in reader]
    rows: list[dict[str, str]] = []
    for line in content.splitlines():
        parts = [part.strip() for part in re.split(r"[,，\t]", line) if part.strip()]
        if len(parts) >= 3:
            rows.append({"subject": parts[0], "relation": parts[1], "object": parts[2]})
    return rows


def map_rows_to_triples(rows: list[dict[str, str]], evidence_id: str | None = None, filename: str | None = None) -> list[ExtractedTriple]:
    triples: list[ExtractedTriple] = []
    for index, row in enumerate(rows, start=1):
        lowered = {key.strip().lower(): value.strip() for key, value in row.items()}
        if _is_case_timeline_row(row):
            triples.extend(_timeline_row_to_triples(row, lowered, index, evidence_id))
            continue
        if _is_call_record_row(row):
            triples.extend(_call_row_to_triples(row, lowered, index, evidence_id, filename))
            continue

        is_bank_flow = _is_bank_flow_row(row)
        graph_eligible = _bank_flow_graph_eligible(row, lowered, filename) if is_bank_flow else True

        subject = _pick(row, lowered, "subject", "主体", "账户真实姓名", "账户假名", "付款方", "交易方", "姓名", "用户ID") or f"记录{index}"
        relation = _pick(row, lowered, "relation", "关系", "借贷类型", "资金方向_按账户", "交易业务类型", "交易用途类型") or "关联"
        obj = _pick(row, lowered, "object", "客体", "对手真实姓名", "对手假名", "对手方ID", "收款方", "对方户名", "value") or "未识别对象"
        amount = _pick(row, lowered, "交易金额_元", "金额", "交易金额", "amount")
        time = _pick(row, lowered, "映射案情时间", "交易时间", "原始交易时间", "time", "date")
        properties = {
            "row_index": index,
            "amount": amount,
            "time": time,
            "original_time": _pick(row, lowered, "原始交易时间"),
            "mapped_case_time": _pick(row, lowered, "映射案情时间"),
            "time_mapping_rule": _pick(row, lowered, "时间映射规则"),
            "time_mapping_note": _pick(row, lowered, "时间映射说明"),
            "account_alias": _pick(row, lowered, "账户假名"),
            "counterparty_alias": _pick(row, lowered, "对手假名"),
            "is_bank_flow": is_bank_flow,
            "is_case_person": _pick(row, lowered, "是否案件映射人物"),
            "is_core_counterparty": _pick(row, lowered, "是否案件核心对手方"),
            "source_file": _pick(row, lowered, "来源文件"),
            "source_dataset": Path(filename).name if filename else None,
            "graph_eligible": graph_eligible,
        }
        props = {key: value for key, value in properties.items() if value not in (None, "")}
        triples.append(ExtractedTriple(subject=subject, relation=relation, object=obj, evidence_id=evidence_id, properties=props))
        if amount:
            triples.append(ExtractedTriple(subject=subject, relation="交易金额", object=amount, evidence_id=evidence_id, properties=props))
        if time:
            triples.append(ExtractedTriple(subject=subject, relation="发生时间", object=time, evidence_id=evidence_id, properties=props))
    return triples


def _timeline_row_to_triples(
    row: dict[str, str],
    lowered: dict[str, str],
    index: int,
    evidence_id: str | None,
) -> list[ExtractedTriple]:
    fact = _pick(row, lowered, "案件事实") or f"案件事实{index}"
    people = _pick(row, lowered, "相关人") or "案件时间线"
    time = _pick(row, lowered, "时间")
    evidence = _pick(row, lowered, "相关证据")
    properties = {
        "row_index": index,
        "time": time,
        "mapped_case_time": time,
        "case_fact": fact,
        "related_people": people,
        "related_evidence": evidence,
        "source_dataset": "case_timeline",
        "graph_eligible": True,
    }
    props = {key: value for key, value in properties.items() if value not in (None, "")}
    triples = [ExtractedTriple(subject=people, relation="案件事实", object=fact, evidence_id=evidence_id, properties=props)]
    if evidence:
        triples.append(ExtractedTriple(subject=fact, relation="对应证据", object=evidence, evidence_id=evidence_id, properties=props))
    return triples


def _is_case_timeline_row(row: dict[str, str]) -> bool:
    return any("案件事实" in key for key in row)



def _is_call_record_row(row: dict[str, str]) -> bool:
    keys = "".join(row.keys()).lower()
    source = str(row.get("\u6765\u6e90\u6587\u4ef6") or "").lower()
    tokens = (
        "\u4e3b\u53eb",
        "\u88ab\u53eb",
        "\u901a\u8bdd",
        "\u77ed\u4fe1",
        "\u624b\u673a\u53f7",
        "\u53f7\u7801",
        "call",
        "phone",
        "caller",
        "callee",
        "cdr",
    )
    return any(token in keys for token in tokens) or "call_records" in source or "cdr" in source


def _call_row_to_triples(
    row: dict[str, str],
    lowered: dict[str, str],
    index: int,
    evidence_id: str | None,
    filename: str | None,
) -> list[ExtractedTriple]:
    subject = _pick(
        row,
        lowered,
        "\u4e3b\u53eb\u771f\u5b9e\u59d3\u540d",
        "\u53d1\u9001\u4eba\u771f\u5b9e\u59d3\u540d",
        "\u4e3b\u53eb\u5047\u540d",
        "\u53d1\u9001\u4eba",
        "caller_name",
        "caller",
    ) or f"\u901a\u8bdd\u4e3b\u4f53{index}"
    obj = _pick(
        row,
        lowered,
        "\u88ab\u53eb\u771f\u5b9e\u59d3\u540d",
        "\u63a5\u6536\u4eba\u771f\u5b9e\u59d3\u540d",
        "\u88ab\u53eb\u5047\u540d",
        "\u63a5\u6536\u4eba",
        "callee_name",
        "callee",
    ) or "\u672a\u8bc6\u522b\u901a\u8bdd\u5bf9\u8c61"
    call_type = _pick(row, lowered, "\u901a\u8bdd\u7c7b\u578b", "\u77ed\u4fe1\u7c7b\u578b", "\u7c7b\u578b", "call_type") or "\u901a\u8bdd"
    relation = "\u77ed\u4fe1\u8054\u7cfb" if "\u77ed\u4fe1" in call_type else "\u901a\u8bdd"
    time = _pick(row, lowered, "\u6620\u5c04\u6848\u60c5\u65f6\u95f4", "\u901a\u8bdd\u5f00\u59cb\u65f6\u95f4", "\u77ed\u4fe1\u65f6\u95f4", "\u53d1\u9001\u65f6\u95f4", "time", "date")
    properties = {
        "row_index": index,
        "time": time,
        "mapped_case_time": _pick(row, lowered, "\u6620\u5c04\u6848\u60c5\u65f6\u95f4"),
        "original_time": _pick(row, lowered, "\u901a\u8bdd\u5f00\u59cb\u65f6\u95f4", "\u77ed\u4fe1\u65f6\u95f4", "\u53d1\u9001\u65f6\u95f4"),
        "duration_seconds": _pick(row, lowered, "\u65f6\u957f_\u79d2", "\u901a\u8bdd\u65f6\u957f", "duration"),
        "call_id": _pick(row, lowered, "\u901a\u8bddID", "\u77ed\u4fe1ID", "call_id", "id"),
        "caller_number": _pick(row, lowered, "\u4e3b\u53eb\u53f7\u7801", "\u53d1\u9001\u53f7\u7801", "caller_number"),
        "callee_number": _pick(row, lowered, "\u88ab\u53eb\u53f7\u7801", "\u63a5\u6536\u53f7\u7801", "callee_number"),
        "caller_alias": _pick(row, lowered, "\u4e3b\u53eb\u5047\u540d", "\u53d1\u9001\u4eba\u5047\u540d"),
        "callee_alias": _pick(row, lowered, "\u88ab\u53eb\u5047\u540d", "\u63a5\u6536\u4eba\u5047\u540d"),
        "event_stage": _pick(row, lowered, "\u4e8b\u4ef6\u9636\u6bb5"),
        "chain_type": _pick(row, lowered, "\u94fe\u6761\u7c7b\u578b"),
        "evidence_relevance": _pick(row, lowered, "\u8bc1\u636e\u76f8\u5173\u6027"),
        "source_file": _pick(row, lowered, "\u6765\u6e90\u6587\u4ef6"),
        "source_dataset": Path(filename).name if filename else None,
        "graph_eligible": True,
        "is_call_record": True,
        "subject_type": "person",
        "object_type": "person",
        "relation_category": "communication",
        "tags": "call_record",
    }
    props = {key: value for key, value in properties.items() if value not in (None, "")}
    return [ExtractedTriple(subject=subject, relation=relation, object=obj, evidence_id=evidence_id, properties=props)]


def _is_bank_flow_row(row: dict[str, str]) -> bool:
    if _is_call_record_row(row):
        return False
    keys = "".join(row.keys())
    return any(
        token in keys
        for token in (
            "\u4ea4\u6613\u91d1\u989d",
            "\u8d26\u6237\u771f\u5b9e\u59d3\u540d",
            "\u5bf9\u624b\u771f\u5b9e\u59d3\u540d",
            "\u8d44\u91d1\u65b9\u5411",
            "\u501f\u8d37\u7c7b\u578b",
            "amount",
        )
    )


def _bank_flow_graph_eligible(row: dict[str, str], lowered: dict[str, str], filename: str | None) -> bool:
    dataset = (filename or "").lower()
    if "case_relevant_bank_flows" in dataset:
        return True
    if "cleaned_bank_flows" in dataset:
        return _truthy(_pick(row, lowered, "是否案件核心对手方"))
    return _truthy(_pick(row, lowered, "是否案件核心对手方")) or _truthy(_pick(row, lowered, "是否案件映射人物"))


def _truthy(value: str | None) -> bool:
    return str(value or "").strip().lower() in {"true", "1", "yes", "y", "是", "核心"}


def build_passages_from_triples(triples: list[ExtractedTriple]) -> list[PassageRecord]:
    passages: list[PassageRecord] = []
    for index, triple in enumerate(triples):
        extra = []
        if triple.properties.get("amount"):
            extra.append(f"金额 {triple.properties['amount']}")
        if triple.properties.get("time"):
            extra.append(f"时间 {triple.properties['time']}")
        suffix = f"（{'，'.join(extra)}）" if extra else ""
        passages.append(PassageRecord(text=f"{triple.subject} --{triple.relation}-> {triple.object}{suffix}", evidence_id=triple.evidence_id, triple_index=index))
    return passages


def extract_unstructured_triples(content: str, evidence_id: str | None = None) -> tuple[list[ExtractedTriple], list[PassageRecord]]:
    sentences = [item.strip() for item in re.split(r"[。！？；\?\!\n]+", content) if item.strip()]
    triples: list[ExtractedTriple] = []
    passages: list[PassageRecord] = []
    for index, sentence in enumerate(sentences[:80], start=1):
        # 抽取时间
        time_value = _extract_time_from_text(sentence)
        
        people = re.findall(r"[\u4e00-\u9fa5]{2,4}", sentence)
        subject = people[0] if people else f"文本片段{index}"
        obj = people[1] if len(people) > 1 else sentence[:80]
        
        # 构建属性，包含时间
        properties = {"source_dataset": "rule_stub", "graph_eligible": False}
        if time_value:
            properties["time"] = time_value
        
        triple = ExtractedTriple(
            subject=subject,
            relation=_infer_relation(sentence),
            object=obj,
            evidence_id=evidence_id,
            properties=properties,
        )
        triples.append(triple)
        passages.append(PassageRecord(text=sentence, evidence_id=evidence_id, triple_index=index - 1))
        
        # 如果有时间，生成独立的时间三元组
        if time_value and subject != f"文本片段{index}":
            time_triple = ExtractedTriple(
                subject=subject,
                relation="发生时间",
                object=time_value,
                evidence_id=evidence_id,
                properties={"source_dataset": "rule_stub", "graph_eligible": True, "time": time_value},
            )
            triples.append(time_triple)
    
    if not triples and content.strip():
        triples.append(
            ExtractedTriple(
                subject="文本证据",
                relation="包含",
                object=content[:120],
                evidence_id=evidence_id,
                properties={"source_dataset": "rule_stub", "graph_eligible": False},
            )
        )
        passages.append(PassageRecord(text=content.strip(), evidence_id=evidence_id, triple_index=0))
    return triples, passages
def _extract_time_from_text(text: str) -> str | None:
    """从文本中抽取时间信息，返回标准化的时间字符串。"""
    # 匹配 YYYY年M月D日 格式
    match = re.search(r'((?:19|20)\d{2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日', text)
    if match:
        return f"{match.group(1)}年{match.group(2)}月{match.group(3)}日"
    
    # 匹配 YYYY年M月 格式
    match = re.search(r'((?:19|20)\d{2})\s*年\s*(\d{1,2})\s*月', text)
    if match:
        return f"{match.group(1)}年{match.group(2)}月"
    
    # 匹配 YYYY年 格式
    match = re.search(r'((?:19|20)\d{2})\s*年', text)
    if match:
        return f"{match.group(1)}年"
    
    # 匹配 YYYY-MM-DD 格式
    match = re.search(r'((?:19|20)\d{2})-(\d{1,2})-(\d{1,2})', text)
    if match:
        return f"{match.group(1)}年{match.group(2)}月{match.group(3)}日"
    
    # 匹配 YYYY/MM/DD 格式
    match = re.search(r'((?:19|20)\d{2})/(\d{1,2})/(\d{1,2})', text)
    if match:
        return f"{match.group(1)}年{match.group(2)}月{match.group(3)}日"
    
    return None

def _infer_relation(sentence: str) -> str:
    if any(word in sentence for word in ("转账", "收款", "付款", "流水", "金额")):
        return "资金往来"
    if any(word in sentence for word in ("明知", "故意", "徇私", "隐瞒")):
        return "主观状态线索"
    if any(word in sentence for word in ("职务", "审批", "办理", "释放", "拘留", "立案")):
        return "职务行为"
    return "提及"


def _pick(original: dict[str, str], lowered: dict[str, str], *names: str) -> str | None:
    for name in names:
        key = name.lower()
        if key in lowered and lowered[key]:
            return lowered[key]
        for source_key, value in original.items():
            if name in source_key and value:
                return value.strip()
    return None


def _cell(value: Any) -> str:
    return "" if value is None else str(value).strip()
